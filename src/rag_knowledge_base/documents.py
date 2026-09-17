"""Document parsers for PDF, DOCX, Markdown, and text files."""

from __future__ import annotations

import csv
import io
from pathlib import Path

from .chunking import normalize_text
from .exceptions import DocumentValidationError
from .models import LoadedPage


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt", ".csv"}


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentValidationError(f"无法识别文本编码: {path.name}")


def _parse_pdf(path: Path) -> list[LoadedPage]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentValidationError(
            "解析 PDF 需要安装 pypdf，请执行 pip install -e '.[rag]'"
        ) from exc

    reader = PdfReader(str(path))
    pages: list[LoadedPage] = []
    for index, page in enumerate(reader.pages, start=1):
        text = normalize_text(page.extract_text() or "")
        if text:
            pages.append(LoadedPage(text=text, page=index, metadata={"page": index}))
    if not pages:
        raise DocumentValidationError("PDF 中没有可提取的文本，扫描版 PDF 请先做 OCR")
    return pages


def _parse_docx(path: Path) -> list[LoadedPage]:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentValidationError(
            "解析 DOCX 需要安装 python-docx，请执行 pip install -e '.[rag]'"
        ) from exc

    document = Document(str(path))
    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                blocks.append(" | ".join(values))
    text = normalize_text("\n\n".join(blocks))
    if not text:
        raise DocumentValidationError("DOCX 中没有可提取的文本")
    return [LoadedPage(text=text, page=1, metadata={"page": None})]


def _parse_csv(path: Path) -> list[LoadedPage]:
    text = _read_text(path)
    rows: list[str] = []
    reader = csv.reader(io.StringIO(text))
    for row in reader:
        cleaned = [cell.strip() for cell in row]
        if any(cleaned):
            rows.append(" | ".join(cleaned))
    content = normalize_text("\n".join(rows))
    if not content:
        raise DocumentValidationError("CSV 中没有可提取的数据")
    return [LoadedPage(text=content, page=1, metadata={"page": None})]


def parse_document(path: Path) -> list[LoadedPage]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentValidationError(f"不支持的文件类型 {suffix}，支持: {supported}")

    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix == ".csv":
        return _parse_csv(path)

    text = normalize_text(_read_text(path))
    if not text:
        raise DocumentValidationError("文档内容为空")
    return [LoadedPage(text=text, page=1, metadata={"page": None})]
